from datetime import datetime, timezone
from pathlib import Path
from tempfile import TemporaryDirectory

import pytest
from openpyxl import load_workbook

from radar_desk.data.database import Database
from radar_desk.data.report_repository import ReportRepository
from radar_desk.data.ticket_repository import TicketRepository
from radar_desk.models.report import ReportFilter
from radar_desk.models.ticket import TicketInput
from radar_desk.services.report_service import ReportService, ReportValidationError
from radar_desk.services.ticket_service import TicketService


def ticket_input(**changes: str) -> TicketInput:
    values = {
        "title": "Falha de acesso ao sistema",
        "description": "O usuário não consegue acessar o sistema interno.",
        "category": "Acesso e permissões",
        "priority": "Alta",
        "status": "Aberto",
        "assignee": "Equipe N1",
    }
    values.update(changes)
    return TicketInput(**values)


def test_report_indicators_period_and_rating() -> None:
    with TemporaryDirectory() as directory:
        database = Database(Path(directory) / "test.sqlite3")
        database.initialize()
        current = datetime(2026, 8, 1, 8, 0, tzinfo=timezone.utc)
        ticket_service = TicketService(
            TicketRepository(database), clock=lambda: current
        )
        report_service = ReportService(
            ReportRepository(database), clock=lambda: current
        )

        resolved = ticket_service.create(ticket_input())
        current = datetime(2026, 8, 1, 11, 0, tzinfo=timezone.utc)
        ticket_service.update(resolved.id, ticket_input(status="Resolvido"))
        current = datetime(2026, 8, 2, 8, 0, tzinfo=timezone.utc)
        active = ticket_service.create(
            ticket_input(
                title="Servidor indisponível",
                description="O servidor principal não responde ao monitoramento.",
                category="Rede e internet",
                priority="Crítica",
            )
        )
        current = datetime(2026, 8, 2, 20, 0, tzinfo=timezone.utc)

        snapshot = report_service.build()

        assert snapshot.total == 2
        assert snapshot.active == 1
        assert snapshot.completed == 1
        assert snapshot.overdue == 1
        assert snapshot.overdue_tickets == (active,)
        assert snapshot.average_resolution_hours == pytest.approx(3.0)
        assert dict(snapshot.status_counts)["Resolvido"] == 1
        assert dict(snapshot.category_counts)["Rede e internet"] == 1
        assert snapshot.volume_by_day == (("2026-08-01", 1), ("2026-08-02", 1))

        rating = report_service.rate_ticket(resolved.id, 5, "Atendimento rápido.")
        rated = report_service.build(ReportFilter(end_date="2026-08-01"))

        assert rating.rating == 5
        assert rated.total == 1
        assert rated.average_rating == pytest.approx(5.0)
        assert rated.rating_count == 1
        assert report_service.repository.get_rating(resolved.id).comment == "Atendimento rápido."

        with pytest.raises(ReportValidationError):
            report_service.rate_ticket(active.id, 4)
        with pytest.raises(ReportValidationError):
            report_service.build(
                ReportFilter(start_date="2026-08-10", end_date="2026-08-01")
            )


def test_report_exports_valid_excel_and_pdf_files() -> None:
    with TemporaryDirectory() as directory:
        root = Path(directory)
        database = Database(root / "test.sqlite3")
        database.initialize()
        now = datetime(2026, 8, 10, 14, 0, tzinfo=timezone.utc)
        ticket_service = TicketService(TicketRepository(database), clock=lambda: now)
        ticket_service.create(ticket_input(title="Exportar indicadores"))
        report_service = ReportService(
            ReportRepository(database), clock=lambda: now
        )
        snapshot = report_service.build()

        excel_path = report_service.export_excel(snapshot, root / "relatorio.xlsx")
        pdf_path = report_service.export_pdf(snapshot, root / "relatorio.pdf")

        workbook = load_workbook(excel_path, read_only=True)
        assert workbook.sheetnames == [
            "Resumo",
            "Indicadores",
            "Volume por período",
            "Chamados",
        ]
        assert workbook["Chamados"]["A2"].value == "RD-2026-0001"
        workbook.close()
        assert pdf_path.read_bytes().startswith(b"%PDF")
        assert pdf_path.stat().st_size > 1_000
