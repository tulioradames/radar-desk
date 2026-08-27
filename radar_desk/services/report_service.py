"""Indicadores, avaliações e exportações da versão 0.8."""

from __future__ import annotations

from collections import Counter
from collections.abc import Callable
from datetime import datetime
from pathlib import Path

from radar_desk.data.report_repository import ReportRepository
from radar_desk.models.report import ReportFilter, ReportSnapshot, TicketRating
from radar_desk.models.ticket import CATEGORIES, PRIORITIES, STATUSES
from radar_desk.services.automation_service import AutomationService
from radar_desk.services.change_tracker import ChangeTracker


class ReportValidationError(ValueError):
    pass


class ReportService:
    def __init__(
        self,
        repository: ReportRepository,
        clock: Callable[[], datetime] | None = None,
        automation: AutomationService | None = None,
        tracker: ChangeTracker | None = None,
    ) -> None:
        self.repository = repository
        self.clock = clock or (lambda: datetime.now().astimezone())
        self.automation = automation or AutomationService(self.clock)
        self.tracker = tracker

    def build(self, filters: ReportFilter | None = None) -> ReportSnapshot:
        filters = filters or ReportFilter()
        if filters.start_date and filters.end_date and filters.start_date > filters.end_date:
            raise ReportValidationError("A data inicial não pode ser posterior à data final.")

        tickets = self.repository.list_tickets(filters)
        status_counter = Counter(ticket.status for ticket in tickets)
        category_counter = Counter(ticket.category for ticket in tickets)
        priority_counter = Counter(ticket.priority for ticket in tickets)
        volume_counter = Counter(ticket.created_at[:10] for ticket in tickets)
        active = sum(
            status_counter[status]
            for status in ("Aberto", "Em andamento", "Aguardando")
        )
        completed = status_counter["Resolvido"] + status_counter["Encerrado"]
        overdue_tickets = tuple(
            ticket
            for ticket in tickets
            if self.automation.sla_for(ticket).is_overdue
        )
        resolved_at = self.repository.resolution_datetimes(
            [ticket.id for ticket in tickets]
        )
        resolution_hours: list[float] = []
        for ticket in tickets:
            resolved_value = resolved_at.get(ticket.id)
            if not resolved_value:
                continue
            opened = datetime.fromisoformat(ticket.created_at)
            resolved = datetime.fromisoformat(resolved_value)
            resolution_hours.append(max(0.0, (resolved - opened).total_seconds() / 3600))

        average_rating, rating_count = self.repository.rating_summary(
            [ticket.id for ticket in tickets]
        )
        generated = self.clock().isoformat(timespec="seconds")
        return ReportSnapshot(
            generated_at=generated,
            start_date=filters.start_date,
            end_date=filters.end_date,
            total=len(tickets),
            active=active,
            completed=completed,
            overdue=len(overdue_tickets),
            average_resolution_hours=(
                sum(resolution_hours) / len(resolution_hours)
                if resolution_hours
                else None
            ),
            average_rating=average_rating,
            rating_count=rating_count,
            status_counts=tuple((status, status_counter[status]) for status in STATUSES),
            category_counts=tuple(
                (category, category_counter[category]) for category in CATEGORIES
            ),
            priority_counts=tuple(
                (priority, priority_counter[priority]) for priority in PRIORITIES
            ),
            volume_by_day=tuple(sorted(volume_counter.items())),
            overdue_tickets=overdue_tickets,
            tickets=tuple(tickets),
        )

    def rate_ticket(self, ticket_id: int, rating: int, comment: str = "") -> TicketRating:
        if rating not in range(1, 6):
            raise ReportValidationError("A avaliação deve estar entre 1 e 5 estrelas.")
        normalized_comment = comment.strip()
        if len(normalized_comment) > 500:
            raise ReportValidationError("O comentário da avaliação deve ter até 500 caracteres.")
        ticket = self.repository.ticket_repository.get(ticket_id)
        if ticket.status not in ("Resolvido", "Encerrado"):
            raise ReportValidationError(
                "Somente chamados resolvidos ou encerrados podem ser avaliados."
            )
        result = self.repository.save_rating(
            ticket_id, rating, normalized_comment, self.clock()
        )
        if self.tracker:
            self.tracker.record(
                "Atendimento avaliado",
                "ticket_rating",
                ticket.protocol,
                "update",
                {
                    "ticket_id": result.ticket_id,
                    "protocol": ticket.protocol,
                    "rating": result.rating,
                    "comment": result.comment,
                    "created_at": result.created_at,
                    "updated_at": result.updated_at,
                },
            )
        return result

    def export_excel(self, snapshot: ReportSnapshot, destination: Path | str) -> Path:
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Alignment, Font, PatternFill
        except ImportError as error:  # pragma: no cover - mensagem para ambiente incompleto
            raise RuntimeError("Instale a dependência openpyxl para exportar Excel.") from error

        path = self._destination(destination, ".xlsx")
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Resumo"
        summary.append(["Radar Desk — Relatório operacional"])
        summary.append(["Gerado em", self._display_datetime(snapshot.generated_at)])
        summary.append(["Período", self._period_label(snapshot)])
        summary.append([])
        summary.append(["Indicador", "Valor"])
        summary.append(["Total de chamados", snapshot.total])
        summary.append(["Chamados ativos", snapshot.active])
        summary.append(["Chamados concluídos", snapshot.completed])
        summary.append(["Chamados atrasados", snapshot.overdue])
        summary.append([
            "Tempo médio de atendimento (horas)",
            round(snapshot.average_resolution_hours, 2)
            if snapshot.average_resolution_hours is not None
            else "Sem dados",
        ])
        summary.append([
            "Avaliação média",
            round(snapshot.average_rating, 2)
            if snapshot.average_rating is not None
            else "Sem avaliações",
        ])
        summary.append(["Total de avaliações", snapshot.rating_count])
        self._style_sheet(summary, PatternFill, Font, Alignment)

        breakdown = workbook.create_sheet("Indicadores")
        breakdown.append(["Status", "Quantidade", "Categoria", "Quantidade", "Prioridade", "Quantidade"])
        rows = max(
            len(snapshot.status_counts),
            len(snapshot.category_counts),
            len(snapshot.priority_counts),
        )
        for index in range(rows):
            values: list[str | int] = []
            for collection in (
                snapshot.status_counts,
                snapshot.category_counts,
                snapshot.priority_counts,
            ):
                values.extend(collection[index] if index < len(collection) else ("", ""))
            breakdown.append(values)
        self._style_sheet(breakdown, PatternFill, Font, Alignment)

        volume = workbook.create_sheet("Volume por período")
        volume.append(["Data", "Chamados abertos"])
        for day, total in snapshot.volume_by_day:
            volume.append([self._display_date(day), total])
        self._style_sheet(volume, PatternFill, Font, Alignment)

        tickets_sheet = workbook.create_sheet("Chamados")
        tickets_sheet.append([
            "Protocolo", "Título", "Categoria", "Prioridade", "Status",
            "Responsável", "Abertura", "Atualização",
        ])
        for ticket in snapshot.tickets:
            tickets_sheet.append([
                ticket.protocol,
                ticket.title,
                ticket.category,
                ticket.priority,
                ticket.status,
                ticket.assignee or "Não atribuído",
                self._display_datetime(ticket.created_at),
                self._display_datetime(ticket.updated_at),
            ])
        self._style_sheet(tickets_sheet, PatternFill, Font, Alignment)

        path.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(path)
        return path

    def export_pdf(self, snapshot: ReportSnapshot, destination: Path | str) -> Path:
        try:
            from reportlab.lib import colors
            from reportlab.lib.enums import TA_CENTER
            from reportlab.lib.pagesizes import A4, landscape
            from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
            from reportlab.lib.units import mm
            from reportlab.platypus import (
                PageBreak,
                Paragraph,
                SimpleDocTemplate,
                Spacer,
                Table,
                TableStyle,
            )
        except ImportError as error:  # pragma: no cover - mensagem para ambiente incompleto
            raise RuntimeError("Instale a dependência reportlab para exportar PDF.") from error

        path = self._destination(destination, ".pdf")
        path.parent.mkdir(parents=True, exist_ok=True)
        styles = getSampleStyleSheet()
        title_style = ParagraphStyle(
            "RadarTitle",
            parent=styles["Title"],
            textColor=colors.HexColor("#0f766e"),
            alignment=TA_CENTER,
            spaceAfter=10,
        )
        heading_style = ParagraphStyle(
            "RadarHeading",
            parent=styles["Heading2"],
            textColor=colors.HexColor("#172033"),
            spaceBefore=8,
            spaceAfter=7,
        )
        document = SimpleDocTemplate(
            str(path),
            pagesize=landscape(A4),
            rightMargin=14 * mm,
            leftMargin=14 * mm,
            topMargin=14 * mm,
            bottomMargin=14 * mm,
            title="Radar Desk — Relatório operacional",
            author="Radar Desk",
        )
        story = [
            Paragraph("Radar Desk — Relatório operacional", title_style),
            Paragraph(
                f"Gerado em {self._display_datetime(snapshot.generated_at)} · "
                f"{self._period_label(snapshot)}",
                styles["Normal"],
            ),
            Spacer(1, 8),
        ]
        indicator_data = [
            ["Total", "Ativos", "Concluídos", "Atrasados", "Tempo médio", "Avaliação"],
            [
                str(snapshot.total),
                str(snapshot.active),
                str(snapshot.completed),
                str(snapshot.overdue),
                self._hours_label(snapshot.average_resolution_hours),
                self._rating_label(snapshot.average_rating, snapshot.rating_count),
            ],
        ]
        story.append(self._pdf_table(indicator_data, Table, TableStyle, colors))
        story.append(Paragraph("Chamados por status", heading_style))
        story.append(
            self._pdf_table(
                [["Status", "Quantidade"], *[[name, str(total)] for name, total in snapshot.status_counts]],
                Table,
                TableStyle,
                colors,
            )
        )
        story.append(Paragraph("Chamados por categoria", heading_style))
        story.append(
            self._pdf_table(
                [["Categoria", "Quantidade"], *[[name, str(total)] for name, total in snapshot.category_counts]],
                Table,
                TableStyle,
                colors,
            )
        )
        story.append(PageBreak())
        story.append(Paragraph("Detalhamento dos chamados", heading_style))
        ticket_rows = [["Protocolo", "Título", "Categoria", "Prioridade", "Status", "Responsável", "Abertura"]]
        for ticket in snapshot.tickets:
            ticket_rows.append([
                ticket.protocol,
                ticket.title[:42],
                ticket.category,
                ticket.priority,
                ticket.status,
                (ticket.assignee or "Não atribuído")[:24],
                self._display_datetime(ticket.created_at),
            ])
        if len(ticket_rows) == 1:
            ticket_rows.append(["—", "Nenhum chamado no período", "—", "—", "—", "—", "—"])
        story.append(self._pdf_table(ticket_rows, Table, TableStyle, colors, repeat_rows=1))
        document.build(story)
        return path

    @staticmethod
    def _destination(destination: Path | str, suffix: str) -> Path:
        path = Path(destination)
        return path if path.suffix.lower() == suffix else path.with_suffix(suffix)

    @staticmethod
    def _style_sheet(sheet, PatternFill, Font, Alignment) -> None:
        header_fill = PatternFill("solid", fgColor="0F766E")
        header_font = Font(color="FFFFFF", bold=True)
        for row in sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)
        header_rows = [1] if sheet.title != "Resumo" else [1, 5]
        for row_number in header_rows:
            for cell in sheet[row_number]:
                if cell.value is not None:
                    cell.fill = header_fill
                    cell.font = header_font
        for column in sheet.columns:
            letter = column[0].column_letter
            width = min(48, max(12, max(len(str(cell.value or "")) for cell in column) + 2))
            sheet.column_dimensions[letter].width = width
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions

    @staticmethod
    def _pdf_table(data, Table, TableStyle, colors, repeat_rows: int = 1):
        table = Table(data, repeatRows=repeat_rows, hAlign="LEFT")
        table.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#0f766e")),
                    ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
                    ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"),
                    ("FONTNAME", (0, 1), (-1, -1), "Helvetica"),
                    ("FONTSIZE", (0, 0), (-1, -1), 8),
                    ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#cbd5e1")),
                    ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#f8fafc")]),
                    ("VALIGN", (0, 0), (-1, -1), "TOP"),
                    ("LEFTPADDING", (0, 0), (-1, -1), 5),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 5),
                    ("TOPPADDING", (0, 0), (-1, -1), 5),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
                ]
            )
        )
        return table

    @staticmethod
    def _display_date(value: str) -> str:
        try:
            return datetime.fromisoformat(value).strftime("%d/%m/%Y")
        except ValueError:
            return value

    @classmethod
    def _display_datetime(cls, value: str) -> str:
        try:
            return datetime.fromisoformat(value).strftime("%d/%m/%Y %H:%M")
        except ValueError:
            return value

    @classmethod
    def _period_label(cls, snapshot: ReportSnapshot) -> str:
        if snapshot.start_date and snapshot.end_date:
            return f"Período: {cls._display_date(snapshot.start_date)} a {cls._display_date(snapshot.end_date)}"
        if snapshot.start_date:
            return f"Desde {cls._display_date(snapshot.start_date)}"
        if snapshot.end_date:
            return f"Até {cls._display_date(snapshot.end_date)}"
        return "Todos os períodos"

    @staticmethod
    def _hours_label(value: float | None) -> str:
        return f"{value:.1f} h" if value is not None else "Sem dados"

    @staticmethod
    def _rating_label(value: float | None, count: int) -> str:
        return f"{value:.1f}/5 ({count})" if value is not None else "Sem avaliações"
